import pandas as pd
import matplotlib.pyplot as plt
from docx import Document
from docx.shared import Inches
import io
import numpy as np

# === Load data ===
data_path = "nb_analytics_small.xlsx"
xls = pd.ExcelFile(data_path, engine="openpyxl")

df_sales = pd.read_excel(xls, "nb_fact_sales")
df_log = pd.read_excel(xls, "nb_logistics")
df_dist = pd.read_excel(xls, "nb_distributor_kpis")
df_geo = pd.read_excel(xls, "nb_geography")

# Ensure revenue column
if "revenue" not in df_sales.columns:
    df_sales["revenue"] = df_sales["qty"] * df_sales.get("effective_price", df_sales.get("list_price"))

# === Insights Preparation ===

# 1. Top states revenue
state_rev = df_sales.groupby("state", as_index=False).agg(revenue=("revenue","sum"))
top_states = state_rev.sort_values("revenue", ascending=False).head(5)

# 2. Worst lanes by cost per km
df_log["cost_per_km"] = df_log["freight_cost_ngn"] / df_log["distance_km"].replace(0, np.nan)
lane_summary = df_log.groupby(["plant","state"], as_index=False).agg(
    avg_distance=("distance_km","mean"),
    avg_lead=("lead_time_days","mean"),
    on_time=("on_time_ratio","mean"),
    cost_per_km=("cost_per_km","mean"),
    breakage=("breakage_ratio","mean")
)
worst_lanes_cost = lane_summary.sort_values("cost_per_km", ascending=False).head(5)

# 3. Distributor risk
dist = df_dist.copy()
dist["credit_utilization_pct"] = (
    (dist["credit_utilization"]*100)
    if dist["credit_utilization"].max() <= 1.2
    else (dist["credit_utilization"]/dist["credit_utilization"].max()*100)
)
dist["risk_score"] = (
    0.15
    + 0.25*(dist["credit_utilization_pct"]>85)
    + 0.2*(dist["dso_days"]>45)
    - 0.2*(dist["otif_ratio"]>0.9)
)
dist["risk_score"] = (dist["risk_score"] + dist["churn_risk"].fillna(0)).clip(0,1)
dist_risk = dist.groupby("distributor", as_index=False).agg(
    dso=("dso_days","mean"),
    otif=("otif_ratio","mean"),
    credit=("credit_utilization_pct","mean"),
    churn=("churn_risk","mean"),
    risk=("risk_score","mean")
).sort_values("risk", ascending=False).head(5)

# === Generate Charts ===

# Top states revenue bar
fig1, ax1 = plt.subplots()
ax1.bar(top_states["state"], top_states["revenue"]/1e9)
ax1.set_ylabel("Revenue (₦ Billion)")
ax1.set_title("Top 5 States by Revenue")
img_buf1 = io.BytesIO()
plt.savefig(img_buf1, format='png', bbox_inches="tight")
img_buf1.seek(0)
plt.close(fig1)

# Worst lanes cost bar
fig2, ax2 = plt.subplots()
ax2.barh(worst_lanes_cost["plant"] + "→" + worst_lanes_cost["state"], worst_lanes_cost["cost_per_km"])
ax2.set_xlabel("₦ per km")
ax2.set_title("Worst 5 Lanes by Cost per km")
img_buf2 = io.BytesIO()
plt.savefig(img_buf2, format='png', bbox_inches="tight")
img_buf2.seek(0)
plt.close(fig2)

# Distributor risk bar
fig3, ax3 = plt.subplots()
ax3.bar(dist_risk["distributor"], dist_risk["risk"])
ax3.set_ylabel("Risk Score (0-1)")
ax3.set_title("Top 5 At-Risk Distributors")
img_buf3 = io.BytesIO()
plt.savefig(img_buf3, format='png', bbox_inches="tight")
img_buf3.seek(0)
plt.close(fig3)

# === Create Word doc with visuals ===
docx_path = "NB_CEO_Insights_Report_With_Charts.docx"
document = Document()

document.add_heading("Nigerian Breweries – CEO/Directors Data-Driven Insights", level=0)
document.add_paragraph(
    "This document summarizes key insights from the nb_analytics_small.xlsx dataset, "
    "with embedded visuals for CEO/Directors review."
)

# Section 1: Top states
document.add_heading("Top States by Revenue", level=1)
document.add_paragraph("The chart below highlights the top 5 states contributing the highest revenues.")
document.add_picture(img_buf1, width=Inches(5))

# Section 2: Worst logistics lanes
document.add_heading("Worst Lanes by Cost-to-Serve", level=1)
document.add_paragraph("The chart below shows the five most expensive logistics lanes (₦ per km).")
document.add_picture(img_buf2, width=Inches(5))

# Section 3: Distributor Risk
document.add_heading("Distributor Risk – Top 5", level=1)
document.add_paragraph(
    "The chart below shows the 5 distributors with highest risk scores, "
    "considering credit utilization, DSO, OTIF, and churn risk."
)
document.add_picture(img_buf3, width=Inches(5))

document.add_page_break()
document.add_paragraph(
    "This report was auto-generated to support executive-level decision-making "
    "on Nigerian Breweries' sales, logistics, and distributor management."
)

document.save(docx_path)

print(f"✅ Report generated: {docx_path}")